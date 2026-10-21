"""Branded download files for the Data page, written to web/downloads/.
- netra-dataset.xlsx         cover sheet (about, sources, licence, citation) + one sheet per table
- netra-companies.csv/.json  every reviewed company group
- netra-reported-emissions.csv  company-reported emissions with the exact source
- README.txt
Reads web/data/{companies,hidden,index,factors}.json (from build.py and summaries.py) and data/reported_emissions.csv.
Usage: python pipeline/exports.py
"""
import csv
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "web" / "data"
OUT = ROOT / "web" / "downloads"
URL = "https://netra.rofihosted.space"
TODAY = date.today().isoformat()
INK, PAPER, SIGNAL = "121212", "FBFAF7", "E8590C"
GRADES = [(0.6, "A"), (0.4, "B"), (0.2, "C"), (0.1, "D"), (0, "E")]

ABOUT = {
    "name": "Netra",
    "tagline": "See what a carbon claim is really worth.",
    "url": URL,
    "generated": TODAY,
    "citation": f"Netra (2026). Carbon claim dataset. IEEE ClimateChain Global Hackathon 2026. {URL}",
    "licence": "Data derived from the Berkeley Carbon Trading Project VROD is shared under CC BY 4.0; "
               "keep the attribution to Berkeley and to Netra. Code: MIT.",
    "sources": [
        "Berkeley Carbon Trading Project, Voluntary Registry Offsets Database v2026-08 (CC BY 4.0)",
        "CarbonPlan OffsetsDB (project types)",
        "ICVCM assessment status, snapshot 18 Aug 2026",
        "Probst et al. 2024, Nature Communications, doi:10.1038/s41467-024-53645-z",
        "Climate TRACE v5.7.0 (ownership-weighted emissions, 2025)",
        "Company annual and sustainability reports (one URL per row)",
    ],
    "method": "Likely real tonnes = tonnes x the share of real reductions measured for the project type in "
              "peer-reviewed research; only where three sources agree on the type. Grade A >= 60%, B 40-60%, "
              "C 20-40%, D 10-20%, E < 10% of researched tonnes. Averages per project type, not verdicts on any project.",
    "methods_document": "https://github.com/rofiperlungoding/climatechain-id/blob/main/docs/METHODS.md",
}


def grade(e, k, t):
    if not k or k / t < 0.25:
        return ""
    return next(g for m, g in GRADES if e / k >= m)


def company_rows(companies):
    out = []
    for c in companies:
        r, gap, ct = c["rep"], c.get("gap"), c["ct"]
        out.append({
            "company_group": c["group"],
            "registry_names": "; ".join(c["names"]),
            "retired_t": c["t"],
            "researched_t": c["k"],
            "likely_real_t": c["e"],
            "likely_real_low_t": c["el"],
            "likely_real_high_t": c["eh"],
            "likely_real_share_of_researched": round(c["e"] / c["k"], 4) if c["k"] else "",
            "grade": grade(c["e"], c["k"], c["t"]),
            "ccp_label_t": c["ccp"],
            "reported_year": r["year"] if r else "",
            "reported_scope1_t": (r["scope1"] or "") if r else "",
            "reported_total_t": (r["total"] or "") if r else "",
            "reported_source": r["source"] if r else "",
            "reported_url": r["url"] if r else "",
            "no_report_reason": gap["reason"] if gap else "",
            "independent_low_t": ct["low"] if ct else "",
            "independent_high_t": ct["high"] if ct else "",
            "independent_year": ct["year"] if ct else "",
            "netra_page": f"{URL}/index.html#buyer-{c['i']}",
        })
    return out


def write_csv(path, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:  # BOM so Excel opens UTF-8 names correctly
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    cols = list(rows[0])
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True, color=PAPER)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"
    for i, c in enumerate(cols, 1):
        width = max(len(str(c)), *(len(str(r[c])) for r in rows[:200]))
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
        if c.endswith("_t"):
            for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for x in cell:
                    x.number_format = "#,##0"
        if c.endswith("_share_of_researched") or c == "share":
            for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for x in cell:
                    x.number_format = "0.0%"
    return ws


def cover(wb, meta, n_companies):
    ws = wb.active
    ws.title = "About Netra"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 100
    ws["B2"] = "netra"
    ws["B2"].font = Font(size=28, bold=True, color=INK)
    ws["B3"] = ABOUT["tagline"]
    ws["B3"].font = Font(size=14, color=SIGNAL, bold=True)
    lines = [
        ("Website", URL), ("Generated", TODAY),
        ("What this is", f"Carbon credits retired on Verra and Gold Standard since 2020 ({meta['tonnes']:,} t), "
                         f"checked against peer-reviewed research, company reports and independent estimates. "
                         f"{n_companies} reviewed company groups."),
        ("Method", ABOUT["method"]),
        ("Sheets", "Companies · Hidden buyers · Project types · Years · Reported emissions"),
        ("Cite as", ABOUT["citation"]), ("Licence", ABOUT["licence"]),
        ("Full method", ABOUT["methods_document"]),
    ] + [("Source" if i == 0 else "", s) for i, s in enumerate(ABOUT["sources"])]
    row = 5
    for k, v in lines:
        ws.cell(row, 2, k).font = Font(bold=True, color="55524C")
        c = ws.cell(row, 3, v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if v.startswith("http"):
            c.hyperlink = v
            c.font = Font(color=SIGNAL, underline="single")
        row += 1


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    companies = json.loads((DATA / "companies.json").read_text(encoding="utf-8"))
    hidden = json.loads((DATA / "hidden.json").read_text(encoding="utf-8"))
    index = json.loads((DATA / "index.json").read_text(encoding="utf-8"))
    factors = json.loads((DATA / "factors.json").read_text(encoding="utf-8"))["factors"]
    meta = index["meta"]
    reported = list(csv.DictReader(open(ROOT / "data" / "reported_emissions.csv", encoding="utf-8")))

    comp = company_rows(companies)
    hid = [{"kind": k, "retired_t": v["t"], "share": round(v["t"] / hidden["total"], 4), "accounts": v["buyers"]}
           for k, v in hidden["kinds"].items()]
    types = [{"project_type": factors[c]["label"], "retired_t": t, "share": round(t / meta["tonnes"], 4),
              "research_share_real": factors[c]["oar"] if factors[c]["oar"] is not None else "no estimate",
              "low": factors[c]["lo"] if factors[c]["oar"] is not None else "",
              "high": factors[c]["hi"] if factors[c]["oar"] is not None else "",
              "source": factors[c].get("source", "")} for c, t in sorted(meta["cats"].items(), key=lambda x: -x[1])]
    years = [{"year": int(y), "retired_t": v[0], "researched_t": v[2], "likely_real_t": v[1],
              "likely_real_share_of_researched": round(v[1] / v[2], 4) if v[2] else ""} for y, v in sorted(meta["ys"].items())]
    rep = [{"company_group": r["group"], "year": r["year"], "scope1_t": int(r["scope1_t"]) if r["scope1_t"] else "",
            "total_t": int(r["total_t"]) if r["total_t"] else "", "boundary": r["boundary"], "source": r["source_title"],
            "url": r["source_url"], "section": r["section"], "checked": r["checked"]} for r in reported]

    wb = Workbook()
    cover(wb, meta, len(comp))
    sheet(wb, "Companies", comp)
    sheet(wb, "Hidden buyers", hid)
    sheet(wb, "Project types", types)
    sheet(wb, "Years", years)
    sheet(wb, "Reported emissions", rep)
    wb.properties.creator = "Netra"
    wb.properties.title = "Netra carbon claim dataset"
    wb.save(OUT / "netra-dataset.xlsx")

    write_csv(OUT / "netra-companies.csv", comp)
    write_csv(OUT / "netra-reported-emissions.csv", rep)
    (OUT / "netra-companies.json").write_text(json.dumps({"netra": ABOUT, "companies": comp}, ensure_ascii=False, indent=1), encoding="utf-8")
    (OUT / "README.txt").write_text(
        f"NETRA · {ABOUT['tagline']}\n{URL}\nGenerated {TODAY}\n\n"
        "Files\n  netra-dataset.xlsx            everything, with an About sheet\n"
        "  netra-companies.csv / .json   one row per reviewed company group\n"
        "  netra-reported-emissions.csv  company-reported emissions, with the exact source\n\n"
        f"Method\n  {ABOUT['method']}\n\nSources\n" + "".join(f"  - {s}\n" for s in ABOUT["sources"]) +
        f"\nCite as\n  {ABOUT['citation']}\n\nLicence\n  {ABOUT['licence']}\n", encoding="utf-8")
    assert len(comp) == len(companies) and all(r["netra_page"].startswith(URL) for r in comp)
    print(f"wrote {len(comp)} companies, {len(rep)} reported rows to {OUT}")


if __name__ == "__main__":
    main()
